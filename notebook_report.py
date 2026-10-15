"""Export the saved notebook, without changing or executing it, to PDF and DOCX.

Both writers consume the same ordered content, including code and rich outputs.
The PDF uses vector equations; DOCX uses editable Office Math equations.
"""

from __future__ import annotations

import base64
from copy import deepcopy
from dataclasses import dataclass
from functools import lru_cache
from html import escape
from io import BytesIO
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from latex2mathml.converter import convert
from lxml import etree, html
import markdown
import matplotlib
from PIL import Image as PILImage
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Flowable, Image, PageBreak, Paragraph, Preformatted, SimpleDocTemplate,
    Spacer, Table, TableStyle,
)

ROOT = Path(__file__).resolve().parent
NOTEBOOK = ROOT / "Assignment3.ipynb"
TITLE = "Assignment 3"
AUTHOR = "Albert Kabore"
WIDTH = 6 * inch
MATH_NS = "http://www.w3.org/1998/Math/MathML"
OFFICE_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def source(value) -> str:
    return value if isinstance(value, str) else "".join(value)


@dataclass
class Block:
    kind: str
    value: object
    cell: int
    label: str = ""


@lru_cache(maxsize=None)
def mathml(expression: str):
    root = etree.fromstring(convert(expression.strip()).encode("utf-8"))
    # latex2mathml emits matrix brackets as siblings. Make them fences so Office
    # Math stretches the brackets over the full matrix height.
    for parent in list(root.iter()):
        i = 0
        while i + 2 < len(parent):
            left, matrix, right = parent[i:i + 3]
            if (etree.QName(left).localname == "mo" and left.text == "["
                    and etree.QName(matrix).localname == "mtable"
                    and etree.QName(right).localname == "mo" and right.text == "]"):
                fenced = etree.Element(f"{{{MATH_NS}}}mfenced", open="[", close="]")
                parent.remove(matrix)
                fenced.append(matrix)
                parent.remove(right)
                parent.replace(left, fenced)
            i += 1
    return root


def content(notebook: dict) -> tuple[list[Block], list[str]]:
    blocks, equations = [], []

    def add_markdown(text: str, cell: int, page_images: bool = False,
                     image_options: dict | None = None):
        def replace(match):
            equations.append(match.group(2).strip())
            index = len(equations) - 1
            if match.group(1) == "$$":
                return f'\n\n<div data-equation="{index}"></div>\n\n'
            return f'<span data-equation="{index}"></span>'

        text = re.sub(r"(\${1,2})(.+?)\1", replace, text, flags=re.S)
        rendered = markdown.markdown(text, extensions=["tables", "fenced_code"])
        for element in html.fragments_fromstring(rendered):
            if isinstance(element, str):
                raise ValueError(f"Unparsed text in notebook cell {cell}")
            if (element.tag == "p" and len(element) == 1
                    and element[0].tag == "img"):
                picture = element[0]
                path = (ROOT / picture.get("src", "")).resolve()
                if not path.is_relative_to(ROOT) or not path.is_file():
                    raise ValueError(f"Invalid notebook image path: {path}")
                data = path.read_bytes()
                if image_options:
                    with PILImage.open(BytesIO(data)) as original:
                        mode = "L" if image_options.get("grayscale") else "RGB"
                        resized = original.convert(mode)
                        width = min(resized.width, int(image_options["max_width"]))
                        height = round(resized.height * width / resized.width)
                        resized = resized.resize((width, height), PILImage.Resampling.LANCZOS)
                        buffer = BytesIO()
                        resized.save(buffer, format="JPEG", optimize=True,
                                     quality=int(image_options["jpeg_quality"]))
                        data = buffer.getvalue()
                blocks.append(Block("page_image" if page_images else "image",
                                    data, cell, picture.get("alt", "")))
                continue
            blocks.append(Block("html", element, cell))

    for index, cell in enumerate(notebook["cells"]):
        if cell["cell_type"] == "markdown":
            add_markdown(source(cell["source"]), index,
                         cell.get("metadata", {}).get("report_page_images", False),
                         cell.get("metadata", {}).get("report_image_options"))
        elif cell["cell_type"] == "code":
            blocks.append(Block("code", source(cell["source"]), index,
                                f'In [{cell.get("execution_count") or " "}]'))
            for output in cell.get("outputs", []):
                if output["output_type"] == "error":
                    raise ValueError(f"Notebook cell {index} has an error output")
                if output["output_type"] == "stream":
                    blocks.append(Block("output", source(output["text"]), index))
                    continue
                data = output.get("data", {})
                if "image/png" in data:
                    blocks.append(Block("image", base64.b64decode(source(data["image/png"])), index))
                elif "text/markdown" in data:
                    add_markdown(source(data["text/markdown"]), index)
                elif "text/latex" in data:
                    add_markdown(source(data["text/latex"]), index)
                elif "text/html" in data:
                    for element in html.fragments_fromstring(source(data["text/html"])):
                        blocks.append(Block("html", element, index))
                elif "text/plain" in data:
                    blocks.append(Block("output", source(data["text/plain"]), index))
                elif data:
                    raise ValueError(f"Unsupported output in cell {index}: {list(data)}")
        else:
            raise ValueError(f"Unsupported notebook cell type: {cell['cell_type']}")
    return blocks, equations


# Small MathML compositor. Coordinates are relative to the baseline, in points.
@dataclass
class MathBox:
    width: float
    above: float
    below: float
    operations: list


def moved(operations, x, y):
    result = []
    for operation in operations:
        kind, *args = operation
        if kind == "text":
            px, py, font, size, text = args
            result.append((kind, px + x, py + y, font, size, text))
        elif kind == "line":
            x1, y1, x2, y2 = args
            result.append((kind, x1 + x, y1 + y, x2 + x, y2 + y))
        elif kind == "rect":
            px, py, width, height = args
            result.append((kind, px + x, py + y, width, height))
    return result


def math_layout(node, size=10.5) -> MathBox:
    tag = etree.QName(node).localname
    children = list(node)
    if tag in {"mi", "mn", "mo", "mtext"}:
        text = node.text or ""
        font = "Eq-I" if tag == "mi" and text.isascii() else "Eq"
        pad = size * .15 if text in {"=", "+", "−", "⊗", "·"} else 0
        return MathBox(pdfmetrics.stringWidth(text, font, size) + 2 * pad,
                       size * .78, size * .22, [("text", pad, 0, font, size, text)])
    if tag == "mspace":
        amount = node.get("width", "0em")
        return MathBox(float(amount.removesuffix("em")) * size, 0, 0, [])
    if tag in {"math", "mrow", "mtd", "mstyle"}:
        parts = [math_layout(child, size) for child in children]
        if not parts:
            return MathBox(0, 0, 0, [])
        operations, x = [], 0
        for part in parts:
            operations.extend(moved(part.operations, x, 0))
            x += part.width
        return MathBox(x, max(p.above for p in parts), max(p.below for p in parts), operations)
    if tag == "mfenced":
        inside = math_layout(children[0], size)
        pad, cap = size * .5, size * .28
        top, bottom = inside.above + 2, -inside.below - 2
        width = inside.width + 2 * pad
        ops = moved(inside.operations, pad, 0)
        for x, direction in [(1, 1), (width - 1, -1)]:
            ops += [("line", x, bottom, x, top),
                    ("line", x, bottom, x + direction * cap, bottom),
                    ("line", x, top, x + direction * cap, top)]
        return MathBox(width, top, -bottom, ops)
    if tag == "mtable":
        rows = [[math_layout(cell, size) for cell in row] for row in children]
        columns = max(map(len, rows))
        widths = [max(row[c].width for row in rows if c < len(row)) for c in range(columns)]
        heights = [max(cell.above for cell in row) + max(cell.below for cell in row) for row in rows]
        gap_x, gap_y = size * 1.05, size * .5
        height = sum(heights) + gap_y * (len(rows) - 1)
        top = height / 2 + size * .24
        cursor, ops = top, []
        for row, row_height in zip(rows, heights):
            baseline = cursor - max(cell.above for cell in row)
            x = 0
            for cell, width in zip(row, widths):
                ops.extend(moved(cell.operations, x + (width - cell.width) / 2, baseline))
                x += width + gap_x
            cursor -= row_height + gap_y
        return MathBox(sum(widths) + gap_x * (columns - 1), top, height - top, ops)
    if tag == "mfrac":
        numerator, denominator = [math_layout(child, size * .86) for child in children]
        width = max(numerator.width, denominator.width) + size * .45
        axis, gap = size * .24, size * .18
        ny, dy = axis + gap + numerator.below, axis - gap - denominator.above
        ops = moved(numerator.operations, (width - numerator.width) / 2, ny)
        ops += moved(denominator.operations, (width - denominator.width) / 2, dy)
        ops.append(("line", 0, axis, width, axis))
        return MathBox(width, ny + numerator.above, denominator.below - dy, ops)
    if tag == "msqrt":
        inside = math_layout(etree.Element(f"{{{MATH_NS}}}mrow"), size)
        if children:
            group = etree.Element(f"{{{MATH_NS}}}mrow")
            group.extend(deepcopy(children))
            inside = math_layout(group, size)
        pad, top = size * .7, inside.above + size * .1
        ops = moved(inside.operations, pad, 0)
        ops += [("line", 0, size * .12, pad * .25, size * .2),
                ("line", pad * .25, size * .2, pad * .48, -inside.below),
                ("line", pad * .48, -inside.below, pad * .87, top),
                ("line", pad * .87, top, pad + inside.width + 1, top)]
        return MathBox(pad + inside.width + 1, top, inside.below, ops)
    if tag in {"msub", "msup"}:
        base = math_layout(children[0], size)
        script = math_layout(children[1], size * .68)
        y = -size * .3 if tag == "msub" else base.above * .68
        ops = base.operations + moved(script.operations, base.width + .4, y)
        return MathBox(base.width + .4 + script.width, max(base.above, y + script.above),
                       max(base.below, script.below - y), ops)
    if tag == "menclose" and node.get("notation") == "box":
        inside = math_layout(children[0], size)
        pad = size * .4
        ops = moved(inside.operations, pad, 0)
        ops.append(("rect", 0, -inside.below - pad, inside.width + 2 * pad,
                    inside.above + inside.below + 2 * pad))
        return MathBox(inside.width + 2 * pad, inside.above + pad, inside.below + pad, ops)
    raise ValueError(f"Unsupported equation element: {tag}")


class Equation(Flowable):
    def __init__(self, expression: str):
        super().__init__()
        self.box = math_layout(mathml(expression))
        self.scale = min(1, WIDTH / max(1, self.box.width))
        self.width = self.box.width * self.scale
        self.height = (self.box.above + self.box.below) * self.scale + 6
        self.spaceAfter = 2
        self.hAlign = "CENTER"

    def draw(self):
        canvas = self.canv
        canvas.saveState()
        canvas.translate(0, 3 + self.box.below * self.scale)
        canvas.scale(self.scale, self.scale)
        canvas.setFillColor(colors.black)
        canvas.setStrokeColor(colors.black)
        canvas.setLineWidth(.65)
        for kind, *args in self.box.operations:
            if kind == "text":
                x, y, font, size, text = args
                canvas.setFont(font, size)
                canvas.drawString(x, y, text)
            elif kind == "line":
                canvas.line(*args)
            else:
                canvas.rect(*args, stroke=1, fill=0)
        canvas.restoreState()


def inline_math(node) -> str:
    tag = etree.QName(node).localname
    children = list(node)
    if tag in {"mi", "mo", "mn", "mtext"}:
        text = node.text or ""
        if text in {"=", "+", "−", "⊗", "·"}:
            text = " " + text + " "
        return escape(text)
    if tag in {"math", "mrow", "mtd"}:
        return "".join(map(inline_math, children))
    if tag in {"msup", "msub"}:
        script = "super" if tag == "msup" else "sub"
        return inline_math(children[0]) + f"<{script}>" + inline_math(children[1]) + f"</{script}>"
    if tag == "msqrt":
        return "√(" + "".join(map(inline_math, children)) + ")"
    if tag == "mfrac":
        return "(" + inline_math(children[0]) + ")/(" + inline_math(children[1]) + ")"
    raise ValueError(f"Unsupported inline equation: {tag}")


def pdf_inline(element, equations) -> str:
    result = escape(element.text or "")
    for child in element:
        if "data-equation" in child.attrib:
            text = inline_math(mathml(equations[int(child.get("data-equation"))]))
        else:
            text = pdf_inline(child, equations)
            if child.tag in {"strong", "b"}:
                text = f"<b>{text}</b>"
            elif child.tag in {"em", "i"}:
                text = f"<i>{text}</i>"
            elif child.tag in {"sub", "sup"}:
                tag = "sub" if child.tag == "sub" else "super"
                text = f"<{tag}>{text}</{tag}>"
            elif child.tag == "code":
                text = f'<font name="Mono">{text}</font>'
            elif child.tag == "a":
                text = f'<a href="{escape(child.get("href"), quote=True)}" color="#174C7E">{text}</a>'
            elif child.tag == "br":
                text = "<br/>"
            elif child.tag not in {"span", "p"}:
                raise ValueError(f"Unsupported inline HTML: {child.tag}")
        result += text + escape(child.tail or "")
    return result


def register_fonts():
    fonts = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    for name, file in [("Body", "DejaVuSans.ttf"), ("Body-B", "DejaVuSans-Bold.ttf"),
                       ("Body-I", "DejaVuSans-Oblique.ttf"), ("Body-BI", "DejaVuSans-BoldOblique.ttf"),
                       ("Mono", "DejaVuSansMono.ttf"), ("Eq", "DejaVuSerif.ttf"),
                       ("Eq-I", "DejaVuSerif-Italic.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(fonts / file)))
    pdfmetrics.registerFontFamily("Body", normal="Body", bold="Body-B", italic="Body-I", boldItalic="Body-BI")


def image_size(data, full_page=False):
    with PILImage.open(BytesIO(data)) as image:
        width, height = image.size
    max_height = 8.1 * inch if full_page else 3.6 * inch
    scale = min(WIDTH / width, max_height / height, .9)
    return width * scale, height * scale


def followed_by_equation(blocks, index):
    if index + 1 >= len(blocks):
        return False
    following = blocks[index + 1]
    return following.kind == "html" and "data-equation" in following.value.attrib


def followed_by_histogram_comment(blocks, index):
    if index + 1 >= len(blocks):
        return False
    following = blocks[index + 1]
    return (following.kind == "html" and
            "".join(following.value.itertext()).startswith("Actual simulator counts:"))


def write_pdf(blocks, equations, path):
    register_fonts()
    # Binary image streams preserve the image data and avoid ASCII85 overhead.
    rl_config.useA85 = False
    body = ParagraphStyle("body", fontName="Body", fontSize=9.5, leading=14, spaceAfter=6)
    heading = ParagraphStyle("heading", parent=body, fontName="Body-B", fontSize=13,
                             leading=18, spaceBefore=8, spaceAfter=7, keepWithNext=True)
    subheading = ParagraphStyle("subheading", parent=heading, fontSize=10.5, leading=14)
    title = ParagraphStyle("title", parent=heading, fontSize=18, leading=23, alignment=1, spaceAfter=12)
    code = ParagraphStyle("code", fontName="Mono", fontSize=7.2, leading=9.6,
                          backColor=colors.HexColor("#F3F4F6"), borderPadding=5, spaceAfter=8)
    output = ParagraphStyle("output", parent=code, backColor=None, fontSize=8, leading=11)
    label = ParagraphStyle("label", parent=body, fontSize=8, textColor=colors.HexColor("#5B6470"),
                           spaceAfter=3, keepWithNext=True)
    reference = ParagraphStyle("reference", parent=body, leftIndent=36,
                               firstLineIndent=-36, leading=19, spaceAfter=7)
    cell_style = ParagraphStyle("tablecell", parent=body, fontSize=8, leading=11, spaceAfter=0)
    story, in_references = [], False
    for block_index, block in enumerate(blocks):
        if block.kind == "code":
            story.append(Paragraph(escape(block.label), label))
            story.append(Preformatted(block.value.rstrip(), code, maxLineLength=98, newLineChars="    "))
        elif block.kind == "output":
            story.append(Preformatted(block.value.rstrip(), output, maxLineLength=88, newLineChars="    "))
        elif block.kind in {"image", "page_image"}:
            width, height = image_size(block.value, block.kind == "page_image")
            image = Image(BytesIO(block.value), width=width, height=height, hAlign="CENTER")
            image.spaceAfter = 8
            image.keepWithNext = followed_by_histogram_comment(blocks, block_index)
            story.append(image)
        else:
            element = block.value
            tag = element.tag
            if tag == "hr":
                story.append(PageBreak())
            elif "data-equation" in element.attrib:
                story.append(Equation(equations[int(element.get("data-equation"))]))
            elif tag.startswith("h") and tag[1:].isdigit():
                text = "".join(element.itertext())
                level = int(tag[1:])
                if level == 2 and text.startswith("Question"):
                    story.append(PageBreak())
                    in_references = False
                if text == "References":
                    story.append(PageBreak())
                    in_references = True
                story.append(Paragraph(pdf_inline(element, equations),
                                       title if level == 1 else heading if level == 2 else subheading))
            elif tag == "p":
                style = reference if in_references else body
                paragraph = Paragraph(pdf_inline(element, equations), style)
                if ((len(element) == 1 and element[0].tag == "strong"
                     and not (element[0].tail or "").strip())
                        or followed_by_equation(blocks, block_index)):
                    paragraph.keepWithNext = True
                story.append(paragraph)
            elif tag in {"ul", "ol"}:
                for index, item in enumerate(element, 1):
                    bullet = "• " if tag == "ul" else f"{index}. "
                    story.append(Paragraph(bullet + pdf_inline(item, equations), body))
            elif tag == "table":
                rows = [[Paragraph(pdf_inline(cell, equations), cell_style) for cell in row]
                        for row in element.xpath(".//tr")]
                widths = [46, WIDTH - 46] if len(rows[0]) == 2 else [WIDTH / len(rows[0])] * len(rows[0])
                table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=0)
                table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF2")),
                    ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#AAB4BF")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]))
                story.extend([table, Spacer(1, 8)])
            else:
                raise ValueError(f"Unsupported block HTML: {tag}")

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Body", 8)
        canvas.setFillColor(colors.HexColor("#66717D"))
        canvas.drawString(.75 * inch, .43 * inch, TITLE)
        canvas.drawRightString(letter[0] - .75 * inch, .43 * inch, f"Page {document.page}")
        canvas.restoreState()

    destination = str(path) if isinstance(path, (str, Path)) else path
    document = SimpleDocTemplate(destination, pagesize=letter, leftMargin=.75 * inch,
                                 rightMargin=.75 * inch, topMargin=.65 * inch,
                                 bottomMargin=.7 * inch, title=TITLE, author=AUTHOR)
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def office_transform():
    for base in [Path("C:/Program Files/Microsoft Office"), Path("C:/Program Files (x86)/Microsoft Office")]:
        if base.exists():
            stylesheet = next(base.rglob("MML2OMML.XSL"), None)
            if stylesheet:
                return etree.XSLT(etree.parse(str(stylesheet)))
    raise FileNotFoundError("Microsoft Office's MML2OMML.XSL is required for editable equations")


def add_hyperlink(paragraph, text, url, bold=False, italic=False):
    from docx.opc.constants import RELATIONSHIP_TYPE
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174C7E")
    properties.append(color)
    for name, enabled in [("b", bold), ("i", italic)]:
        if enabled:
            properties.append(OxmlElement(f"w:{name}"))
    run.append(properties)
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def word_math(expression, transform):
    equation = deepcopy(transform(mathml(expression)).getroot())
    box = math_layout(mathml(expression))
    size = min(10.5, 10.5 * WIDTH / max(1, box.width))
    for run in equation.xpath(".//m:r", namespaces={"m": OFFICE_MATH_NS}):
        properties = OxmlElement("w:rPr")
        font = OxmlElement("w:rFonts")
        font.set(qn("w:ascii"), "Cambria Math")
        font.set(qn("w:hAnsi"), "Cambria Math")
        properties.append(font)
        font_size = OxmlElement("w:sz")
        font_size.set(qn("w:val"), str(round(size * 2)))
        properties.append(font_size)
        run.insert(1 if len(run) and run[0].tag == qn("m:rPr") else 0, properties)
    return equation


def word_inline(paragraph, element, equations, transform, bold=False, italic=False, script=None):
    def text_run(text):
        if text:
            run = paragraph.add_run(text)
            run.bold, run.italic = bold, italic
            if script:
                run.font.subscript = script == "sub"
                run.font.superscript = script == "sup"

    text_run(element.text)
    for child in element:
        if "data-equation" in child.attrib:
            equation = equations[int(child.get("data-equation"))]
            paragraph._p.append(word_math(equation, transform))
        elif child.tag == "a":
            add_hyperlink(paragraph, "".join(child.itertext()), child.get("href"), bold, italic)
        elif child.tag == "br":
            paragraph.add_run().add_break()
        else:
            word_inline(paragraph, child, equations, transform,
                        bold or child.tag in {"strong", "b"}, italic or child.tag in {"em", "i"},
                        child.tag if child.tag in {"sub", "sup"} else script)
        text_run(child.tail)


def write_word(blocks, equations, path):
    transform = office_transform()
    document = Document()
    document.core_properties.title, document.core_properties.author = TITLE, AUTHOR
    document.core_properties.subject = "Complete reproduction of Assignment3.ipynb"
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.75)
    section.top_margin, section.bottom_margin = Inches(.65), Inches(.7)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 19), ("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 10.5)]:
        style = document.styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(size), RGBColor.from_string("172B40")
        style.paragraph_format.keep_with_next = True
    for name, size in [("Notebook Code", 7.5), ("Notebook Output", 8)]:
        style = document.styles.add_style(name, 1)
        style.font.name, style.font.size = "Consolas", Pt(size)
        style.paragraph_format.line_spacing = 1
        style.paragraph_format.space_after = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.add_run(TITLE + "\tPage ").font.size = Pt(8)
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6), WD_TAB_ALIGNMENT.RIGHT)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    in_references = False
    for block_index, block in enumerate(blocks):
        if block.kind in {"code", "output"}:
            if block.kind == "code":
                label = document.add_paragraph(block.label)
                label.paragraph_format.keep_with_next = True
                label.paragraph_format.space_after = Pt(3)
                label.runs[0].font.size = Pt(8)
            paragraph = document.add_paragraph(style="Notebook Code" if block.kind == "code" else "Notebook Output")
            paragraph.add_run(block.value.rstrip())
            if block.kind == "code":
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F3F4F6")
                paragraph._p.get_or_add_pPr().append(shading)
        elif block.kind in {"image", "page_image"}:
            width, height = image_size(block.value, block.kind == "page_image")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = followed_by_histogram_comment(blocks, block_index)
            paragraph.add_run().add_picture(BytesIO(block.value), width=Pt(width), height=Pt(height))
        else:
            element = block.value
            tag = element.tag
            if tag == "hr":
                document.add_page_break()
            elif "data-equation" in element.attrib:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(2)
                office_math = OxmlElement("m:oMathPara")
                office_math.append(word_math(equations[int(element.get("data-equation"))], transform))
                paragraph._p.append(office_math)
            elif tag.startswith("h") and tag[1:].isdigit():
                level = int(tag[1:])
                paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {min(3, level - 1)}")
                word_inline(paragraph, element, equations, transform)
                if level == 1:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if level == 2 and "".join(element.itertext()).startswith("Question"):
                    paragraph.paragraph_format.page_break_before = True
                    in_references = False
                if "".join(element.itertext()) == "References":
                    paragraph.paragraph_format.page_break_before = True
                    in_references = True
            elif tag == "p":
                paragraph = document.add_paragraph()
                word_inline(paragraph, element, equations, transform)
                if in_references:
                    paragraph.paragraph_format.left_indent = Inches(.5)
                    paragraph.paragraph_format.first_line_indent = Inches(-.5)
                    paragraph.paragraph_format.line_spacing = 2
                if ((len(element) == 1 and element[0].tag == "strong"
                     and not (element[0].tail or "").strip())
                        or followed_by_equation(blocks, block_index)):
                    paragraph.paragraph_format.keep_with_next = True
            elif tag in {"ul", "ol"}:
                for item in element:
                    paragraph = document.add_paragraph(style="List Bullet" if tag == "ul" else "List Number")
                    word_inline(paragraph, item, equations, transform)
            elif tag == "table":
                rows = element.xpath(".//tr")
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                table.autofit = False
                widths = [46, WIDTH - 46] if len(rows[0]) == 2 else [WIDTH / len(rows[0])] * len(rows[0])
                for column, width in zip(table.columns, widths):
                    column.width = Pt(width)
                for row_index, row in enumerate(rows):
                    for cell_index, cell in enumerate(row):
                        word_cell = table.cell(row_index, cell_index)
                        word_cell.width = Pt(widths[cell_index])
                        paragraph = word_cell.paragraphs[0]
                        paragraph.paragraph_format.space_after = Pt(3)
                        paragraph.paragraph_format.space_before = Pt(3)
                        paragraph.paragraph_format.keep_with_next = row_index < len(rows) - 1
                        word_inline(paragraph, cell, equations, transform, bold=row_index == 0)
                        for run in paragraph.runs:
                            run.font.size = Pt(8)
                    if row_index == 0:
                        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                    prevent_split = OxmlElement("w:cantSplit")
                    table.rows[row_index]._tr.get_or_add_trPr().append(prevent_split)
                document.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                raise ValueError(f"Unsupported block HTML: {tag}")
    document.save(path)


def build() -> tuple[Path, Path]:
    original = NOTEBOOK.read_bytes()
    notebook = json.loads(original)
    blocks, equations = content(notebook)
    register_fonts()
    pdf_path, word_path = ROOT / "Assignment3.pdf", ROOT / "Assignment3.docx"
    (ROOT / "report_assets").mkdir(exist_ok=True)
    # Stage both files before replacing either final report.
    import tempfile
    with tempfile.TemporaryDirectory(prefix="assignment3-", dir=ROOT / "report_assets") as temporary:
        temporary = Path(temporary)
        write_word(blocks, equations, temporary / word_path.name)
        write_pdf(blocks, equations, temporary / pdf_path.name)
        if NOTEBOOK.read_bytes() != original:
            raise RuntimeError("The source notebook changed during the export")
        (temporary / pdf_path.name).replace(pdf_path)
        try:
            (temporary / word_path.name).replace(word_path)
        except PermissionError:
            pending = ROOT / "report_assets" / "runtime" / "pending_reports"
            pending.mkdir(parents=True, exist_ok=True)
            pending_word = pending / word_path.name
            (temporary / word_path.name).replace(pending_word)
            print(f"Could not replace {word_path.name}; it may be open in Word. "
                  f"The corrected Word report is saved at {pending_word}.")
            word_path = pending_word
    print(f"Reproduced {len(notebook['cells'])} notebook cells, {len(equations)} equations, "
          f"{sum(b.kind == 'image' for b in blocks)} figures, "
          f"{sum(b.kind == 'page_image' for b in blocks)} handwritten pages, and "
          f"{sum(b.kind == 'html' and b.value.tag == 'table' for b in blocks)} tables.")
    return pdf_path, word_path


if __name__ == "__main__":
    for report in build():
        print(f"wrote {report}")
